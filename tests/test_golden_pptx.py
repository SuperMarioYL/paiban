"""端到端 golden 回归 + CLI 冒烟：fixtures/sample-longdoc.md → courseware deck。

期望值全部来自 ``fixtures/golden/expected-props.json``。断言的是写进 .pptx 的 OOXML
属性（缩进 EMU、行距、段前后、对齐、避头尾、三字体槽），不是渲染结果——渲染结果只能
在 WPS 里人工确认（见 ``deck.WPS_UNVERIFIED``）。

LLM 只决定页序，版式是确定性的：所以这里走无 key 的规则切分路径，golden 记录的就是
这条路（``plan_source == "rules"``）。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from typer.testing import CliRunner

from paiban import cli, deck, ooxml, typography

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "fixtures" / "sample-longdoc.md"
GOLDEN = json.loads((ROOT / "fixtures" / "golden" / "expected-props.json").read_text("utf-8"))

_A_RPR = qn("a:rPr")

# GitHub runner 会把 ANSI 序列塞进 typer 的 --help 输出，断言前必须剥掉。
_ANSI = re.compile(r"\x1b\[[0-9;?]*[ -/]*[@-~]")


def plain(result) -> str:
    return _ANSI.sub("", result.output or "")


@pytest.fixture()
def built(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> deck.BuildResult:
    """按 golden 记录的那条路径生成 deck：无 key → 规则切分。"""
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    monkeypatch.delenv("DEEPSEEK_BASE_URL", raising=False)
    return deck.build_deck(
        SOURCE, template=GOLDEN["template"], out=tmp_path / "课件.pptx"
    )


# --------------------------------------------------------------------------- #
# golden：deck 形状
# --------------------------------------------------------------------------- #


def test_golden_deck_shape(built: deck.BuildResult) -> None:
    assert built.plan.source == GOLDEN["plan_source"]
    assert built.plan.fallback_reason  # 降级原因必须给出来，不能被吞掉
    assert built.template == GOLDEN["template"]
    assert built.slide_count == GOLDEN["totals"]["slides"]
    assert len(built.plan.slides) == GOLDEN["totals"]["content_slides"]

    pptx = ooxml.Pptx.open(built.path)
    slides = list(pptx.presentation.slides)
    assert slide_title(slides[0]) == GOLDEN["cover_title"]
    assert [slide_title(slide) for slide in slides[1:]] == GOLDEN["slide_titles"]


def test_golden_totals(built: deck.BuildResult) -> None:
    pptx = ooxml.Pptx.open(built.path)
    bodies = list(pptx.iter_text_bodies())
    paragraphs = [paragraph for body in bodies for paragraph in body.paragraphs]
    runs = [run for paragraph in paragraphs for run in ooxml.paragraph_runs(paragraph)]
    cjk_runs = [run for run in runs if deck.contains_cjk(ooxml.run_text(run))]
    squeezable = [
        paragraph
        for paragraph in paragraphs
        if any(typography.squeeze_flags(paragraph_text(paragraph)))
    ]

    totals = GOLDEN["totals"]
    assert len(bodies) == totals["text_bodies"]
    assert len(paragraphs) == totals["paragraphs"]
    assert len(runs) == totals["runs"]
    assert len(cjk_runs) == totals["cjk_runs"]
    assert len(squeezable) == totals["squeezed_paragraphs"]
    assert count_pictures(pptx.presentation) == totals["pictures"] == 0


# --------------------------------------------------------------------------- #
# golden：逐属性
# --------------------------------------------------------------------------- #


def test_golden_paragraph_attributes(built: deck.BuildResult) -> None:
    """每一段都逐属性对齐 golden：正文一段都不例外。"""
    pptx = ooxml.Pptx.open(built.path)
    checked = 0
    for body in pptx.iter_text_bodies():
        expected = (
            GOLDEN["title_paragraph"]
            if body.role == ooxml.TEXT_ROLE_TITLE
            else GOLDEN["body_paragraph"]
        )
        for paragraph in body.paragraphs:
            layout = ooxml.read_paragraph_layout(paragraph.pPr)
            assert layout.indent_emu == expected["indent_emu"]
            assert layout.line_spacing == expected["line_spacing"]
            assert layout.space_before_pt == expected["space_before_pt"]
            assert layout.space_after_pt == expected["space_after_pt"]
            assert layout.alignment == expected["alignment"]
            assert layout.ea_line_break is expected["ea_line_break"]
            assert layout.hanging_punct is expected["hanging_punct"]
            checked += 1
    assert checked == GOLDEN["totals"]["paragraphs"]


def test_golden_font_slots_and_sizes(built: deck.BuildResult) -> None:
    """每个 run 都写了三字体槽和显式字号，中文不会落到方块字体上。"""
    pptx = ooxml.Pptx.open(built.path)
    expected_slots = ooxml.FontSlots(**GOLDEN["font_slots"])
    checked = 0
    for body in pptx.iter_text_bodies():
        expected_size = (
            GOLDEN["title_paragraph"]["font_size_pt"]
            if body.role == ooxml.TEXT_ROLE_TITLE
            else GOLDEN["body_paragraph"]["font_size_pt"]
        )
        for paragraph in body.paragraphs:
            for run in ooxml.paragraph_runs(paragraph):
                assert ooxml.read_font_slots(run.find(_A_RPR)) == expected_slots
                assert ooxml.run_font_size_pt(run) == expected_size
                checked += 1
    assert checked == GOLDEN["totals"]["runs"]


def test_golden_indent_is_two_characters_at_the_declared_size() -> None:
    """缩进 EMU 就是「2 字符 × 字号」，不是随手写的一个数。"""
    body = GOLDEN["body_paragraph"]
    assert body["indent_emu"] == ooxml.points_to_emu(
        body["first_line_indent_chars"] * body["font_size_pt"]
    )
    assert GOLDEN["title_paragraph"]["indent_emu"] is None


def test_every_squeezable_paragraph_has_negative_character_spacing(
    built: deck.BuildResult,
) -> None:
    pptx = ooxml.Pptx.open(built.path)
    for body in pptx.iter_text_bodies():
        for paragraph in body.paragraphs:
            if not any(typography.squeeze_flags(paragraph_text(paragraph))):
                continue
            spacings = [
                run.find(_A_RPR).get("spc")
                for run in ooxml.paragraph_runs(paragraph)
                if run.find(_A_RPR) is not None
            ]
            assert any(
                spacing is not None and int(spacing) < 0 for spacing in spacings
            ), paragraph_text(paragraph)


def test_build_reports_clean_lint_and_verify_agrees(built: deck.BuildResult) -> None:
    assert built.lint.ok
    assert {finding.level for finding in built.lint.findings} == {deck.LINT_OK}
    assert deck.verify(built.path).ok


def test_build_is_deterministic(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """同一份输入两次输出逐属性一致（plan §4：版式永远是确定性的）。"""
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    first = deck.build_deck(SOURCE, template="courseware", out=tmp_path / "a.pptx")
    second = deck.build_deck(SOURCE, template="courseware", out=tmp_path / "b.pptx")
    assert projection(first.path) == projection(second.path)


# --------------------------------------------------------------------------- #
# 模板族
# --------------------------------------------------------------------------- #


def test_builtin_templates_declare_font_size_explicitly() -> None:
    """模板族必须声明字号：否则首行缩进会落到 FALLBACK_FONT_SIZE_PT。"""
    body_sizes = {}
    for name in deck.BUILTIN_TEMPLATES:
        template = deck.load_template(name)
        assert template.body_style.font_size_pt is not None
        assert template.title_style.font_size_pt is not None
        assert template.body_style.first_line_indent_chars == 2.0
        assert template.title_style.first_line_indent_chars is None
        assert typography.resolve_font(template.body_style.font) is not None
        assert typography.resolve_font(template.title_style.font) is not None
        body_sizes[name] = template.body_style.font_size_pt
    assert len(set(body_sizes.values())) == len(body_sizes)


def test_template_family_type_scale_reaches_the_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """模板族的字号参数真的写进了 .pptx，不只是 JSON 里的摆设。"""
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    sizes = {}
    for name in deck.BUILTIN_TEMPLATES:
        result = deck.build_deck(SOURCE, template=name, out=tmp_path / f"{name}.pptx")
        sizes[name] = body_font_size(result.path)
    assert sizes["courseware"] > sizes["report"]


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #


def test_cli_help_lists_every_command() -> None:
    result = CliRunner().invoke(cli.app, ["--help"])
    output = plain(result)
    assert result.exit_code == 0
    for command in ("build", "verify", "fix"):
        assert command in output
    assert "WPS" in output


def test_cli_build_then_verify(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    out = tmp_path / "课件.pptx"
    runner = CliRunner()

    built = runner.invoke(
        cli.app, ["build", str(SOURCE), "-o", str(out), "--template", "courseware"]
    )
    assert built.exit_code == 0, plain(built)
    assert out.is_file()
    assert "rules" in plain(built)
    assert str(out) in plain(built)

    verified = runner.invoke(cli.app, ["verify", str(out)])
    assert verified.exit_code == 0, plain(verified)
    assert deck.LINT_OK in plain(verified)


def test_cli_verify_warns_on_a_deck_without_the_pass(tmp_path: Path) -> None:
    """没跑过版式 pass 的 deck：verify 判 warn 并以退出码 1 结束。"""
    path = tmp_path / "raw.pptx"
    raw_deck().save(path)

    result = CliRunner().invoke(cli.app, ["verify", str(path)])

    assert result.exit_code == 1
    assert deck.LINT_WARN in plain(result)


def test_cli_fix_rewrites_an_existing_deck(tmp_path: Path) -> None:
    path = tmp_path / "raw.pptx"
    raw_deck().save(path)

    result = CliRunner().invoke(
        cli.app, ["fix", str(path), "--template", "courseware"]
    )

    assert result.exit_code == 0, plain(result)
    assert deck.verify(path, template="courseware").ok


def test_cli_rejects_unknown_template(tmp_path: Path) -> None:
    result = CliRunner().invoke(
        cli.app,
        ["build", str(SOURCE), "-o", str(tmp_path / "x.pptx"), "--template", "nope"],
    )
    assert result.exit_code == 2
    assert "courseware" in plain(result)


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #


def slide_title(slide) -> str:
    title = slide.shapes.title
    return "" if title is None else title.text_frame.text


def paragraph_text(paragraph) -> str:
    return "".join(ooxml.run_text(run) for run in ooxml.paragraph_runs(paragraph))


def count_pictures(presentation) -> int:
    return sum(
        1
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )


def body_font_size(path: Path) -> float:
    """deck 里正文 run 的字号（模板族的 body font_size_pt 应当落在它上面）。"""
    pptx = ooxml.Pptx.open(path)
    sizes = {
        ooxml.run_font_size_pt(run)
        for body in pptx.iter_text_bodies()
        if body.role == ooxml.TEXT_ROLE_BODY
        for paragraph in body.paragraphs
        for run in ooxml.paragraph_runs(paragraph)
    }
    assert len(sizes) == 1, sizes
    return sizes.pop()


def raw_deck() -> ooxml.Pptx:
    """一份没有跑过中文版式 pass 的 deck：字体槽 / 缩进 / 行距都是母版默认。"""
    return ooxml.emit(
        ooxml.DeckIR(
            title="未加工的 deck",
            slides=(
                ooxml.Slide(
                    title="第一页",
                    bullets=(ooxml.Bullet("中文要点，。"), ooxml.Bullet("第二条要点")),
                ),
            ),
        )
    )


def projection(path: Path) -> list:
    """段落属性 + run 属性的快照，用于比较两次构建的结果。"""
    pptx = ooxml.Pptx.open(path)
    snapshot = []
    for body in pptx.iter_text_bodies():
        for paragraph in body.paragraphs:
            snapshot.append(
                (
                    body.role,
                    body.label,
                    ooxml.read_paragraph_layout(paragraph.pPr),
                    [
                        (
                            ooxml.run_text(run),
                            ooxml.run_font_size_pt(run),
                            ooxml.read_font_slots(run.find(_A_RPR)),
                            None if run.find(_A_RPR) is None else run.find(_A_RPR).get("spc"),
                        )
                        for run in ooxml.paragraph_runs(paragraph)
                    ],
                )
            )
    return snapshot
