"""长文档 → deck 流水线，以及 .pptx 的版式复核。

职责边界（mvp_plan.md §3 / §4）：

* ``read_document()``：markdown / txt / docx → 文档结构（标题块 + 段落 + 列表项）；
* ``plan_slides()``：文档结构 → SlidePlan（页序 / 标题 / 要点 / 备注）。默认走 DeepSeek
  语义分段；没有 key、或调用失败时降级为确定性规则切分，并把原因写进 ``SlidePlan``；
* ``load_template()``：内置中文模板族 JSON → 母版版式映射 + 正文 / 标题的版式参数；
* ``build_deck()``：以上三者 + ``ooxml.emit()`` + ``typography.CJKLayoutPass`` → .pptx；
* ``lint()`` / ``verify()``：打开 .pptx 复核版式属性（m3 的 lint 报告）。

LLM 只负责语义分段，版式永远是确定性的（plan §4）：同一份 ``DocOutline`` 两次规则
切分逐属性一致；DeepSeek 只影响页序，不写任何 OOXML 属性。

WPS 边界（plan frontmatter ``schema_unverified: true``）：本仓库没有 WPS 环境，因此
``verify()`` 只复核写进 .pptx 的 OOXML 属性，不渲染、不驱动 WPS 进程、也不断言任何
WPS 版本的实际排版结果。这些断言在 ``WPS_UNVERIFIED`` 与 ``WPS_ROUNDTRIP_CHECKLIST``
里显式列出，必须由人在 WPS 里确认。
"""

from __future__ import annotations

import json
import os
import re
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass, fields
from pathlib import Path
from typing import TYPE_CHECKING, Any

from docx import Document
from openai import OpenAI
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape

from . import ooxml, typography

if TYPE_CHECKING:
    from pptx.oxml.text import CT_RegularTextRun, CT_TextParagraph

# --------------------------------------------------------------------------- #
# 文档结构
# --------------------------------------------------------------------------- #

BLOCK_HEADING = "heading"
BLOCK_BULLET = "bullet"
BLOCK_PARAGRAPH = "paragraph"

# v0.1 的输入面（plan §6：只吃 markdown / docx / txt，不吃 PDF / 图片 / 视频）。
MARKDOWN_SUFFIXES = frozenset({".md", ".markdown", ".txt"})
DOCX_SUFFIXES = frozenset({".docx"})


@dataclass(frozen=True)
class DocBlock:
    """文档里的一个内容块。``level`` 对标题是 1-6，对列表项是缩进层级（0 起）。"""

    kind: str
    text: str
    level: int = 0


@dataclass(frozen=True)
class DocOutline:
    """一份长文档的结构：标题 + 内容块序列。"""

    title: str
    blocks: tuple[DocBlock, ...]


def read_document(path: str | Path) -> DocOutline:
    """读长文档，按后缀分派到 markdown / docx 解析器。"""
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in DOCX_SUFFIXES:
        title, blocks = _read_docx(source)
    elif suffix in MARKDOWN_SUFFIXES:
        title, blocks = _read_markdown(source, plain=suffix == ".txt")
    else:
        raise ValueError(f"不支持的输入格式 {suffix or source.name!r}：v0.1 只吃 markdown / txt / docx")
    return DocOutline(title=title, blocks=blocks)


# --------------------------------------------------------------------------- #
# markdown / txt 解析
# --------------------------------------------------------------------------- #

_ATX_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
# 纯文本（.txt）没有 markdown 语法，用中文常见的编号行当章节标题。
_PLAIN_CHAPTER = re.compile(r"^第[一二三四五六七八九十百零〇0-9]+[讲章节课篇部分]\s*\S")
_PLAIN_SECTION = re.compile(r"^(?:[一二三四五六七八九十]+[、.．]|\d+[、.．])\s*\S")
_PLAIN_HEADING_MAX_CHARS = 40
_BULLET = re.compile(r"^([ \t]*)([-*+]|\d+[.)])[ \t]+(.*)$")
_HORIZONTAL_RULE = re.compile(r"^([-*_])\1{2,}$")
_FENCE = re.compile(r"^\s*(?:```|~~~)")
# 行内标记：deck 里不该出现 markdown 记号，读入时就地去掉。
_INLINE_PATTERNS: tuple[tuple[re.Pattern[str], str], ...] = (
    (re.compile(r"!\[[^\]]*\]\([^)]*\)"), ""),
    (re.compile(r"\[([^\]]*)\]\([^)]*\)"), r"\1"),
    (re.compile(r"`([^`]*)`"), r"\1"),
    (re.compile(r"\*\*([^*]+)\*\*"), r"\1"),
    (re.compile(r"\*([^*]+)\*"), r"\1"),
    (re.compile(r"__([^_]+)__"), r"\1"),
)


def _read_markdown(path: Path, *, plain: bool) -> tuple[str, tuple[DocBlock, ...]]:
    """解析 markdown（``plain=True`` 时按 .txt 处理：额外识别中文编号标题行）。"""
    lines = _strip_frontmatter(path.read_text(encoding="utf-8").splitlines())
    blocks: list[DocBlock] = []
    paragraph: list[str] = []
    in_fence = False
    for raw in lines:
        line = raw.rstrip()
        if _FENCE.match(line):
            _flush_paragraph(blocks, paragraph)
            in_fence = not in_fence
            continue
        if in_fence:
            # 代码块逐行成段，保留缩进，也不做行内标记清理
            if line.strip():
                blocks.append(DocBlock(BLOCK_PARAGRAPH, line))
            continue
        stripped = line.strip()
        if not stripped or _HORIZONTAL_RULE.match(stripped):
            _flush_paragraph(blocks, paragraph)
            continue
        heading = _heading(stripped, plain=plain)
        if heading is not None:
            _flush_paragraph(blocks, paragraph)
            blocks.append(heading)
            continue
        bullet = _bullet(line)
        if bullet is not None:
            _flush_paragraph(blocks, paragraph)
            blocks.append(bullet)
            continue
        paragraph.append(stripped)
    _flush_paragraph(blocks, paragraph)
    if plain and _leading_title_line(blocks):
        # 纯文本稿没有标题语法，开头那行短文本就是标题；消费掉，免得和封面重复
        return blocks[0].text, tuple(blocks[1:])
    return _first_heading_title(blocks) or path.stem, tuple(blocks)


def _leading_title_line(blocks: Sequence[DocBlock]) -> bool:
    return (
        bool(blocks)
        and blocks[0].kind == BLOCK_PARAGRAPH
        and len(blocks[0].text) <= _PLAIN_HEADING_MAX_CHARS
    )


def _heading(text: str, *, plain: bool) -> DocBlock | None:
    match = _ATX_HEADING.match(text)
    if match:
        return DocBlock(BLOCK_HEADING, _strip_inline(match.group(2)).strip(), len(match.group(1)))
    if not plain or len(text) > _PLAIN_HEADING_MAX_CHARS:
        return None
    if _PLAIN_CHAPTER.match(text):
        return DocBlock(BLOCK_HEADING, text, 1)
    if _PLAIN_SECTION.match(text):
        return DocBlock(BLOCK_HEADING, text, 2)
    return None


def _bullet(line: str) -> DocBlock | None:
    match = _BULLET.match(line)
    if match is None:
        return None
    indent = match.group(1).expandtabs(2)
    return DocBlock(BLOCK_BULLET, _strip_inline(match.group(3)).strip(), min(len(indent) // 2, 4))


def _flush_paragraph(blocks: list[DocBlock], lines: list[str]) -> None:
    if not lines:
        return
    text = _strip_inline(_join_lines(lines)).strip()
    lines.clear()
    if text:
        blocks.append(DocBlock(BLOCK_PARAGRAPH, text))


def _join_lines(lines: Sequence[str]) -> str:
    """合并被硬换行拆开的段落：中日韩字符之间不加空格，其余按空格连接。"""
    text = lines[0]
    for line in lines[1:]:
        separator = "" if _is_cjk(text[-1]) and _is_cjk(line[0]) else " "
        text += separator + line
    return text


def _strip_inline(text: str) -> str:
    for pattern, replacement in _INLINE_PATTERNS:
        text = pattern.sub(replacement, text)
    return text


def _strip_frontmatter(lines: list[str]) -> list[str]:
    """去掉文件开头的 YAML frontmatter；没有成对的 ``---`` 就不动。"""
    if not lines or lines[0].strip() != "---":
        return lines
    for index, line in enumerate(lines[1:], start=1):
        if line.strip() == "---":
            return lines[index + 1 :]
    return lines


def _first_heading_title(blocks: Sequence[DocBlock]) -> str:
    for block in blocks:
        if block.kind == BLOCK_HEADING and block.level == 1:
            return block.text
    return ""


# --------------------------------------------------------------------------- #
# docx 解析
# --------------------------------------------------------------------------- #

_DOCX_HEADING = re.compile(r"^(?:heading|标题)\s*([1-9])$", re.IGNORECASE)
_DOCX_TITLE_STYLES = frozenset({"title", "标题"})
_DOCX_LIST = re.compile(r"list|列表", re.IGNORECASE)
_DOCX_LIST_LEVEL = re.compile(r"(\d+)\s*$")


def _read_docx(path: Path) -> tuple[str, tuple[DocBlock, ...]]:
    document = Document(str(path))
    blocks: list[DocBlock] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").strip()
        level = _docx_heading_level(style)
        if level:
            blocks.append(DocBlock(BLOCK_HEADING, text, level))
        elif _DOCX_LIST.search(style):
            match = _DOCX_LIST_LEVEL.search(style)
            level = int(match.group(1)) - 1 if match else 0
            blocks.append(DocBlock(BLOCK_BULLET, text, max(level, 0)))
        else:
            blocks.append(DocBlock(BLOCK_PARAGRAPH, text))
    title = (
        (document.core_properties.title or "").strip()
        or _first_heading_title(blocks)
        or path.stem
    )
    return title, tuple(blocks)


def _docx_heading_level(style: str) -> int:
    if style.casefold() in _DOCX_TITLE_STYLES:
        return 1
    match = _DOCX_HEADING.match(style)
    return int(match.group(1)) if match else 0


# --------------------------------------------------------------------------- #
# 文档结构 → SlidePlan
# --------------------------------------------------------------------------- #

MAX_BULLETS_PER_SLIDE = 6
CONTINUATION_SUFFIX = "（续）"

PLAN_SOURCE_DEEPSEEK = "deepseek"
PLAN_SOURCE_RULES = "rules"

DEFAULT_MODEL = "deepseek-chat"
DEFAULT_BASE_URL = "https://api.deepseek.com"
API_KEY_ENV = "DEEPSEEK_API_KEY"
BASE_URL_ENV = "DEEPSEEK_BASE_URL"

PLAN_SYSTEM_PROMPT = """\
你是中文演示文稿的结构编辑。把用户给的课程稿 / 汇报稿改写成幻灯片页序。

只输出 JSON，形如：
{"slides": [{"title": "页标题", "bullets": ["要点一", "要点二"], "notes": "讲稿"}]}

规则：
- 页标题不超过 20 个汉字；每页 3-5 条要点，每条不超过 40 个汉字
- 要点只提炼文档里已有的信息，不要编造数字、事实、结论或引用
- 背景、例子、口播稿放进 notes，不要塞进要点
- 页数控制在 8-16 页；只做章节分隔的页可以没有 bullets
"""


@dataclass(frozen=True)
class SlidePlan:
    """一份 deck 的页序方案，以及这份方案是怎么来的。

    ``source`` 为 ``deepseek`` 时 ``model`` 是实际用的模型；为 ``rules`` 时
    ``fallback_reason`` 一定非空，调用方必须把它显示给用户——规则切分只保证页序可用，
    不保证语义分段质量。
    """

    slides: tuple[ooxml.Slide, ...]
    source: str
    model: str | None = None
    fallback_reason: str = ""

    def deck(self, title: str, subtitle: str = "") -> ooxml.DeckIR:
        """补上封面标题，得到渲染层要的 DeckIR。"""
        return ooxml.DeckIR(title=title, subtitle=subtitle, slides=self.slides)


def plan_slides(
    outline: DocOutline,
    *,
    client: OpenAI | None = None,
    model: str = DEFAULT_MODEL,
) -> SlidePlan:
    """文档结构 → SlidePlan。

    ``client`` 给定时用它；否则读 ``DEEPSEEK_API_KEY`` 自建 OpenAI 兼容客户端；都没有
    key 就直接规则切分。模型调用失败（网络、鉴权、返回体解析）同样降级为规则切分，
    但原因会写进 ``SlidePlan.fallback_reason``。
    """
    client = client or _default_client()
    if client is None:
        return _rules_plan(outline, f"未设置 {API_KEY_ENV}，按文档结构规则切分")
    try:
        slides = _plan_with_deepseek(outline, client, model)
    except Exception as error:  # LLM 边界：网络 / 鉴权 / 返回体都可能失败，一律降级
        return _rules_plan(outline, f"DeepSeek 调用失败，按文档结构规则切分：{error}")
    return SlidePlan(slides=slides, source=PLAN_SOURCE_DEEPSEEK, model=model)


def _default_client() -> OpenAI | None:
    api_key = os.environ.get(API_KEY_ENV, "").strip()
    if not api_key:
        return None
    base_url = os.environ.get(BASE_URL_ENV, "").strip() or DEFAULT_BASE_URL
    return OpenAI(api_key=api_key, base_url=base_url)


def _rules_plan(outline: DocOutline, reason: str) -> SlidePlan:
    return SlidePlan(
        slides=_plan_with_rules(outline),
        source=PLAN_SOURCE_RULES,
        fallback_reason=reason,
    )


def _plan_with_deepseek(
    outline: DocOutline,
    client: OpenAI,
    model: str,
) -> tuple[ooxml.Slide, ...]:
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": PLAN_SYSTEM_PROMPT},
            {"role": "user", "content": outline_text(outline)},
        ],
        response_format={"type": "json_object"},
        temperature=0.0,
    )
    return _parse_slides(response.choices[0].message.content or "")


def outline_text(outline: DocOutline) -> str:
    """把文档结构还原成带层级的纯文本，作为 DeepSeek 的输入。"""
    lines: list[str] = []
    for block in outline.blocks:
        if block.kind == BLOCK_HEADING:
            lines.append(f"{'#' * min(block.level, 6)} {block.text}")
        elif block.kind == BLOCK_BULLET:
            lines.append(f"{'  ' * min(block.level, 3)}- {block.text}")
        else:
            lines.append(block.text)
    return "\n".join(lines)


def _parse_slides(payload: str) -> tuple[ooxml.Slide, ...]:
    data = json.loads(payload)
    if not isinstance(data, Mapping):
        raise ValueError("模型返回的不是 JSON 对象")
    entries = data.get("slides")
    if not isinstance(entries, list):
        raise ValueError("模型返回的 JSON 里没有 slides 数组")
    slides: list[ooxml.Slide] = []
    for entry in entries:
        if not isinstance(entry, Mapping):
            continue
        raw_bullets = entry.get("bullets")
        items = raw_bullets if isinstance(raw_bullets, list) else ()
        bullets = tuple(ooxml.Bullet(text) for text in map(_as_text, items) if text)
        title = _as_text(entry.get("title"))
        notes = _as_text(entry.get("notes"))
        if not title and not bullets and not notes:
            continue
        slides.append(_slide(title, bullets, notes))
    if not slides:
        raise ValueError("模型返回的 slides 是空的")
    return tuple(slides)


def _as_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


@dataclass
class _Section:
    """一个章节：标题 + 它下面的内容块（构造期可变，切完章节就不再动）。"""

    title: str
    level: int
    blocks: list[DocBlock]


def _sections(outline: DocOutline) -> tuple[_Section, ...]:
    """按标题块切章节；首个标题之前的内容归到文档标题下。"""
    sections: list[_Section] = []
    current: _Section | None = None
    for block in outline.blocks:
        if block.kind == BLOCK_HEADING:
            if current is not None:
                sections.append(current)
            current = _Section(block.text, block.level, [])
        else:
            if current is None:
                current = _Section(outline.title, 1, [])
            current.blocks.append(block)
    if current is not None:
        sections.append(current)
    return tuple(sections)


def _plan_with_rules(outline: DocOutline) -> tuple[ooxml.Slide, ...]:
    slides: list[ooxml.Slide] = []
    for section in _sections(outline):
        bullets, notes = _section_content(section)
        if not bullets and not notes and section.title == outline.title:
            # 文档标题既是封面又是章节：没内容就不用再出一页
            continue
        for index, chunk in enumerate(_chunk(bullets)):
            title = section.title if index == 0 else f"{section.title}{CONTINUATION_SUFFIX}"
            slides.append(_slide(title, chunk, notes if index == 0 else ""))
    return tuple(slides)


def _section_content(section: _Section) -> tuple[tuple[ooxml.Bullet, ...], str]:
    """要点优先取列表项，正文段落进备注；没有列表项时正文段落自己当要点。"""
    listed = [block for block in section.blocks if block.kind == BLOCK_BULLET]
    prose = [block for block in section.blocks if block.kind == BLOCK_PARAGRAPH]
    if listed:
        bullets = tuple(ooxml.Bullet(block.text, min(block.level, 4)) for block in listed)
        return bullets, "\n".join(block.text for block in prose)
    return tuple(ooxml.Bullet(block.text) for block in prose), ""


def _chunk(bullets: Sequence[ooxml.Bullet]) -> Iterator[tuple[ooxml.Bullet, ...]]:
    """按每页上限切要点；没有要点也产出一个空页（章节页）。"""
    for start in range(0, max(len(bullets), 1), MAX_BULLETS_PER_SLIDE):
        yield tuple(bullets[start : start + MAX_BULLETS_PER_SLIDE])


def _slide(title: str, bullets: tuple[ooxml.Bullet, ...], notes: str) -> ooxml.Slide:
    """没有要点的页用章节版式（大标题），有要点的页用正文版式。

    章节版式不保证有正文占位符，所以只在要点为空时才用它——否则要点会被静默丢掉。
    """
    kind = ooxml.SLIDE_CONTENT if bullets else ooxml.SLIDE_SECTION
    return ooxml.Slide(title=title, bullets=bullets, notes=notes, kind=kind)


# --------------------------------------------------------------------------- #
# 内置模板族
# --------------------------------------------------------------------------- #

TEMPLATE_DIR = Path(__file__).parent / "templates"
BUILTIN_TEMPLATES = ("report", "courseware")
DEFAULT_TEMPLATE_NAME = "report"


@dataclass(frozen=True)
class DeckTemplate:
    """一个中文模板族：母版版式映射 + 正文 / 标题的版式参数。"""

    spec: ooxml.TemplateSpec
    body_style: typography.ParagraphStyle
    title_style: typography.ParagraphStyle


def load_template(name: str = DEFAULT_TEMPLATE_NAME) -> DeckTemplate:
    """加载内置模板族。v0.1 只有 ``report``（汇报）与 ``courseware``（课件），不读外部文件。

    模板族 JSON 的 schema::

        {
          "name": "courseware",
          "layouts": {"cover": 0, "content": 1, "section": 2, "end": 5},
          "title": {"font": "微软雅黑", "font_size_pt": 32, "first_line_indent_chars": null,
                    "line_spacing": 1.5, "space_before_pt": 0, "space_after_pt": 6,
                    "alignment": null},
          "body":  {"font": "微软雅黑", "font_size_pt": 18, "first_line_indent_chars": 2,
                    "line_spacing": 1.5, "space_before_pt": 0, "space_after_pt": 6,
                    "alignment": "just"}
        }

    ``layouts`` 的键固定是 ooxml 的页角色常量（cover / content / section / end），值是
    python-pptx 内置母版的 ``slide_layout`` 序号；``title`` / ``body`` 的键就是
    ``typography.ParagraphStyle`` 的字段，缺省项走 ParagraphStyle 的默认值。
    """
    if name not in BUILTIN_TEMPLATES:
        raise ValueError(f"没有内置模板族 {name!r}，可选：{'、'.join(BUILTIN_TEMPLATES)}")
    path = TEMPLATE_DIR / f"{name}.json"
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise ValueError(f"模板族文件缺失：{path}") from None
    return _parse_template(data, name)


def _parse_template(data: Any, name: str) -> DeckTemplate:
    if not isinstance(data, Mapping):
        raise ValueError(f"模板族 {name} 的 JSON 顶层必须是对象")
    layouts = data.get("layouts")
    if not isinstance(layouts, Mapping):
        raise ValueError(f"模板族 {name} 缺 layouts 映射")
    spec = ooxml.TemplateSpec(
        name=str(data.get("name") or name),
        layouts={str(key): int(value) for key, value in layouts.items()},
    )
    body_style = _parse_style(data.get("body"), typography.BODY_STYLE, name, "body")
    title_style = _parse_style(data.get("title"), typography.TITLE_STYLE, name, "title")
    for label, style in (("body", body_style), ("title", title_style)):
        if typography.resolve_font(style.font) is None:
            raise ValueError(f"模板族 {name} 的 {label} 字体 {style.font!r} 不在别名表里")
    return DeckTemplate(spec=spec, body_style=body_style, title_style=title_style)


def _parse_style(
    data: Any,
    default: typography.ParagraphStyle,
    name: str,
    label: str,
) -> typography.ParagraphStyle:
    if data is None:
        return default
    if not isinstance(data, Mapping):
        raise ValueError(f"模板族 {name} 的 {label} 必须是对象")
    known = {field.name for field in fields(typography.ParagraphStyle)}
    unknown = sorted(set(data) - known)
    if unknown:
        raise ValueError(f"模板族 {name} 的 {label} 有不认识的版式参数：{unknown}")
    return typography.ParagraphStyle(
        **{key: data.get(key, getattr(default, key)) for key in known}
    )


# --------------------------------------------------------------------------- #
# 长文档 → .pptx
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class BuildResult:
    """一次 ``build`` 的结果：产物路径、用的页序方案、版式统计、lint 报告。"""

    path: Path
    template: str
    plan: SlidePlan
    layout: typography.LayoutReport
    lint: LintReport

    @property
    def slide_count(self) -> int:
        """封面 + 内容页。"""
        return len(self.plan.slides) + 1


def build_deck(
    source: str | Path,
    *,
    template: str = DEFAULT_TEMPLATE_NAME,
    out: str | Path | None = None,
    client: OpenAI | None = None,
    model: str = DEFAULT_MODEL,
) -> BuildResult:
    """长文档 → 中文模板族 deck，写完 .pptx 后顺手 lint 一次（生成即验收）。

    ``out`` 缺省时与输入同目录同名、换 ``.pptx`` 后缀。
    """
    outline = read_document(source)
    deck_template = load_template(template)
    plan = plan_slides(outline, client=client, model=model)
    pptx = ooxml.emit(plan.deck(outline.title), deck_template.spec)
    layout = typography.CJKLayoutPass(
        body_style=deck_template.body_style,
        title_style=deck_template.title_style,
    ).apply(pptx)
    target = Path(out) if out is not None else Path(source).with_suffix(".pptx")
    pptx.save(target)
    report = lint(
        pptx,
        path=target,
        body_style=deck_template.body_style,
        title_style=deck_template.title_style,
    )
    return BuildResult(
        path=target,
        template=deck_template.spec.name,
        plan=plan,
        layout=layout,
        lint=report,
    )


# --------------------------------------------------------------------------- #
# 版式复核（paiban verify）
# --------------------------------------------------------------------------- #

LINT_OK = "ok"
LINT_WARN = "warn"

VERIFY_SCOPE = "只复核 .pptx 里的 OOXML 属性，不渲染、不驱动 WPS 进程"

# WPS 边界（plan frontmatter ``schema_unverified: true``）：这些结论本仓库无法验证，
# verify 只能把它们列出来，必须由人在 WPS 里确认。
WPS_UNVERIFIED: tuple[str, ...] = (
    "本机是否装了别名表里的字体，以及 WPS 实际选中了哪个字体来渲染中文",
    "标点挤压与避头尾在 WPS 里的渲染结果（行首行尾半宽、禁则处理是否生效）",
    "在 WPS 里改字 / 复制页 / 保存之后版式是否保持（需要人工往返一次）",
)

# 目标版本区间：v0.1 声明支持的 WPS 范围，不是实测结论。
WPS_TARGET_VERSION_RANGE = "WPS Office 2019 及以上（含 WPS 365）"


@dataclass(frozen=True)
class WPSRoundtripStep:
    """WPS 往返验收清单的一步；这五步都要人工在 WPS 里执行。"""

    step: str
    action: str
    expected: str


WPS_ROUNDTRIP_CHECKLIST: tuple[WPSRoundtripStep, ...] = (
    WPSRoundtripStep("打开", "双击 .pptx 用 WPS 打开", "中文字体不回退、没有方块"),
    WPSRoundtripStep("改字", "在任意一页改一段文字", "首行缩进 2 字符与行距 1.5 仍在"),
    WPSRoundtripStep("加页", "复制一页或新建一页", "新版式跟着母版走，缩进行距一致"),
    WPSRoundtripStep("保存", "保存后重新打开", "属性没被 WPS 改写回默认值"),
    WPSRoundtripStep("复核", "对保存后的文件重跑 paiban verify", "告警与保存前一致"),
)

# <a:rPr> 的 ``spc`` 是只读复核用：写侧原语在 ooxml.write_character_spacing。
_A_RPR = qn("a:rPr")


@dataclass(frozen=True)
class LintFinding:
    """一条 lint 结论：``check`` 是稳定标识，``detail`` 是给人看的中文说明。"""

    check: str
    level: str
    detail: str


@dataclass(frozen=True)
class LintReport:
    """``paiban verify`` 的报告：复核了什么、哪里没达标、哪些结论得去 WPS 里看。"""

    slides: int = 0
    text_bodies: int = 0
    paragraphs: int = 0
    runs: int = 0
    findings: tuple[LintFinding, ...] = ()
    path: Path | None = None

    @property
    def ok(self) -> bool:
        return all(finding.level == LINT_OK for finding in self.findings)

    def to_dict(self) -> dict[str, Any]:
        return {
            "path": str(self.path) if self.path is not None else None,
            "scope": VERIFY_SCOPE,
            "slides": self.slides,
            "text_bodies": self.text_bodies,
            "paragraphs": self.paragraphs,
            "runs": self.runs,
            "ok": self.ok,
            "findings": [
                {"check": finding.check, "level": finding.level, "detail": finding.detail}
                for finding in self.findings
            ],
            "wps_target_version_range": WPS_TARGET_VERSION_RANGE,
            "wps_unverified": list(WPS_UNVERIFIED),
        }


@dataclass
class _LintCounts:
    slides: int = 0
    slides_without_placeholder: int = 0
    pictures: int = 0
    text_bodies: int = 0
    paragraphs: int = 0
    runs: int = 0
    cjk_runs: int = 0
    font_resolved: int = 0
    font_unresolved: int = 0
    font_missing_ea: int = 0
    punct_paragraphs: int = 0
    punct_applied: int = 0
    indent_ok: int = 0
    indent_bad: int = 0
    indent_skipped: int = 0
    spacing_ok: int = 0
    spacing_bad: int = 0
    spacing_skipped: int = 0
    align_bad: int = 0
    kinsoku_on: int = 0
    kinsoku_off: int = 0


def verify(path: str | Path, *, template: str | None = None) -> LintReport:
    """打开 .pptx 出 lint 报告。

    ``template`` 给定时按该模板族的版式参数判定期望值；不给就按内置规范（首行缩进 2
    字符、行距 1.5、避头尾开启、无图片兜底）——拿别的工具生成的 .pptx 进来复核时，
    用默认值正是要看的差异。
    """
    deck_template = load_template(template) if template else None
    body_style = deck_template.body_style if deck_template else typography.BODY_STYLE
    title_style = deck_template.title_style if deck_template else typography.TITLE_STYLE
    return lint(
        ooxml.Pptx.open(path),
        path=Path(path),
        body_style=body_style,
        title_style=title_style,
    )


def lint(
    deck: ooxml.Pptx,
    *,
    path: Path | None = None,
    body_style: typography.ParagraphStyle = typography.BODY_STYLE,
    title_style: typography.ParagraphStyle = typography.TITLE_STYLE,
) -> LintReport:
    """复核一个已打开的 deck 的版式属性，期望值来自传入的版式参数。"""
    counts = _LintCounts()
    for slide in deck.presentation.slides:
        counts.slides += 1
        if not any(shape.is_placeholder for shape in slide.shapes):
            counts.slides_without_placeholder += 1
        counts.pictures += _count_pictures(slide.shapes)
    for body in deck.iter_text_bodies():
        counts.text_bodies += 1
        style = title_style if body.role == ooxml.TEXT_ROLE_TITLE else body_style
        for paragraph in body.paragraphs:
            counts.paragraphs += 1
            _lint_paragraph(paragraph, style, counts)
    return LintReport(
        slides=counts.slides,
        text_bodies=counts.text_bodies,
        paragraphs=counts.paragraphs,
        runs=counts.runs,
        findings=_findings(counts),
        path=path,
    )


def _lint_paragraph(
    paragraph: CT_TextParagraph,
    style: typography.ParagraphStyle,
    counts: _LintCounts,
) -> None:
    layout = ooxml.read_paragraph_layout(paragraph.pPr)
    _lint_indent(layout, style, paragraph, counts)
    _lint_spacing(layout, style, counts)
    if style.alignment is not None and layout.alignment != style.alignment:
        counts.align_bad += 1
    if layout.ea_line_break:
        counts.kinsoku_on += 1
    else:
        counts.kinsoku_off += 1
    for run in ooxml.paragraph_runs(paragraph):
        counts.runs += 1
        _lint_run(run, counts)
    _lint_punct(paragraph, counts)


def _lint_indent(
    layout: ooxml.ParagraphLayout,
    style: typography.ParagraphStyle,
    paragraph: CT_TextParagraph,
    counts: _LintCounts,
) -> None:
    if style.first_line_indent_chars is None:
        counts.indent_skipped += 1
        return
    size_pt = _effective_font_size(paragraph, style)
    expected = ooxml.points_to_emu(style.first_line_indent_chars * size_pt)
    if layout.indent_emu is None or abs(layout.indent_emu - expected) > 1:
        counts.indent_bad += 1
    else:
        counts.indent_ok += 1


def _lint_spacing(
    layout: ooxml.ParagraphLayout,
    style: typography.ParagraphStyle,
    counts: _LintCounts,
) -> None:
    if style.line_spacing is None:
        counts.spacing_skipped += 1
    elif layout.line_spacing is None or abs(layout.line_spacing - style.line_spacing) > 1e-6:
        counts.spacing_bad += 1
    else:
        counts.spacing_ok += 1


def _lint_run(run: CT_RegularTextRun, counts: _LintCounts) -> None:
    text = ooxml.run_text(run)
    if not contains_cjk(text):
        return
    counts.cjk_runs += 1
    rPr = run.find(_A_RPR)
    slots = ooxml.read_font_slots(rPr)
    east_asian = slots.ea if slots is not None else None
    if not east_asian:
        counts.font_missing_ea += 1
    elif typography.resolve_font(east_asian) is not None:
        counts.font_resolved += 1
    else:
        counts.font_unresolved += 1


def _lint_punct(paragraph: CT_TextParagraph, counts: _LintCounts) -> None:
    paragraph_text = ooxml.paragraph_text_with_runs(paragraph)
    if not any(typography.squeeze_flags(paragraph_text.text)):
        return
    counts.punct_paragraphs += 1
    if _has_squeezed_run(paragraph):
        counts.punct_applied += 1


def _has_squeezed_run(paragraph: CT_TextParagraph) -> bool:
    for run in ooxml.paragraph_runs(paragraph):
        rPr = run.find(_A_RPR)
        spacing = None if rPr is None else rPr.get("spc")
        if spacing is not None and int(spacing) < 0:
            return True
    return False


def _effective_font_size(
    paragraph: CT_TextParagraph,
    style: typography.ParagraphStyle,
) -> float:
    """段落字号：与 CJKLayoutPass 换算首行缩进时用的是同一条回退链。"""
    return (
        ooxml.paragraph_font_size_pt(paragraph)
        or style.font_size_pt
        or typography.FALLBACK_FONT_SIZE_PT
    )


def _count_pictures(shapes) -> int:
    total = 0
    for shape in shapes:
        if isinstance(shape, GroupShape):
            total += _count_pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            total += 1
    return total


def _findings(counts: _LintCounts) -> tuple[LintFinding, ...]:
    return (
        _font_finding(counts),
        _punct_finding(counts),
        _paragraph_finding(counts),
        _kinsoku_finding(counts),
        _shape_finding(counts),
    )


def _font_finding(counts: _LintCounts) -> LintFinding:
    clean = not counts.font_unresolved and not counts.font_missing_ea
    return LintFinding(
        check="font_alias",
        level=LINT_OK if clean else LINT_WARN,
        detail=(
            f"{counts.cjk_runs} 个中文 run：{counts.font_resolved} 个东亚字体槽已解析，"
            f"{counts.font_unresolved} 个不在别名表，{counts.font_missing_ea} 个没有东亚字体槽"
        ),
    )


def _punct_finding(counts: _LintCounts) -> LintFinding:
    return LintFinding(
        check="punct_compress",
        level=LINT_OK if counts.punct_paragraphs == counts.punct_applied else LINT_WARN,
        detail=(
            f"{counts.punct_paragraphs} 段含连续全角标点，"
            f"其中 {counts.punct_applied} 段已应用挤压（负字距）"
        ),
    )


def _paragraph_finding(counts: _LintCounts) -> LintFinding:
    clean = not counts.indent_bad and not counts.spacing_bad and not counts.align_bad
    return LintFinding(
        check="paragraph_layout",
        level=LINT_OK if clean else LINT_WARN,
        detail=(
            f"首行缩进：{counts.indent_ok} 段达标 / {counts.indent_bad} 段未达标 / "
            f"{counts.indent_skipped} 段不适用；行距：{counts.spacing_ok} 段达标 / "
            f"{counts.spacing_bad} 段未达标；对齐不符 {counts.align_bad} 段"
        ),
    )


def _kinsoku_finding(counts: _LintCounts) -> LintFinding:
    return LintFinding(
        check="kinsoku",
        level=LINT_OK if not counts.kinsoku_off else LINT_WARN,
        detail=(
            f"避头尾（eaLnBrk）：{counts.kinsoku_on} 段已开启，{counts.kinsoku_off} 段未开启"
        ),
    )


def _shape_finding(counts: _LintCounts) -> LintFinding:
    clean = not counts.pictures and not counts.slides_without_placeholder
    return LintFinding(
        check="real_shapes",
        level=LINT_OK if clean else LINT_WARN,
        detail=(
            f"{counts.slides} 页中 {counts.slides - counts.slides_without_placeholder} 页带真实占位符，"
            f"图片 {counts.pictures} 张（Paiban 生成物应为 0）"
        ),
    )


# --------------------------------------------------------------------------- #
# 中文判定
# --------------------------------------------------------------------------- #

_CJK_RANGES = (
    (0x3000, 0x303F),  # 中日韩标点
    (0x3400, 0x4DBF),  # 扩展 A
    (0x4E00, 0x9FFF),  # 基本区
    (0xF900, 0xFAFF),  # 兼容汉字
    (0xFF00, 0xFFEF),  # 全角形式
    (0x20000, 0x2FA1F),  # 扩展 B 及以上
)


def contains_cjk(text: str) -> bool:
    """文本里有没有中日韩字符（含全角标点）；字体槽检查只看这类 run。"""
    return any(_is_cjk(char) for char in text)


def _is_cjk(char: str) -> bool:
    code = ord(char)
    return any(start <= code <= end for start, end in _CJK_RANGES)
