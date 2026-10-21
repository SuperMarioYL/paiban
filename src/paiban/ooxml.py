"""OOXML 读写层：DeckIR → .pptx，以及既有 .pptx 的属性遍历与改写。

职责边界（mvp_plan.md §4）：

* ``emit()``：DeckIR → python-pptx 的真实母版形状与占位符，没有图片兜底；
* ``Pptx.iter_text_bodies()``：打开既有 .pptx，按文本容器 / 段落 / run 粒度
  暴露内容，供 ``paiban fix`` 与 ``paiban verify`` 使用；
* 三字体槽 ``<a:latin>`` / ``<a:ea>`` / ``<a:cs>`` 与段落属性
  （``indent`` / ``lnSpc`` / ``spcBef`` / ``spcAft`` / ``algn`` / ``eaLnBrk`` /
  ``hangingPunct``）的低层读写原语。

版式规则本身（字体别名表、标点挤压、缩进行距）在 typography.py；本模块只提供
机制，不做任何版式决策。本层只读写标准 OOXML，不驱动 WPS 进程、不要求安装
插件、不生成图片。
"""

from __future__ import annotations

import copy
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT
from pptx.oxml.ns import qn
from pptx.oxml.text import (
    CT_RegularTextRun,
    CT_TextCharacterProperties,
    CT_TextParagraph,
    CT_TextParagraphProperties,
)
from pptx.shapes.group import GroupShape
from pptx.util import Emu, Pt

# 文本容器角色：标题类占位符（标题 / 居中标题 / 副标题 / 竖排标题）与其余正文。
TEXT_ROLE_TITLE = "title"
TEXT_ROLE_BODY = "body"

# 页角色：正文页 / 章节页 / 结束页。封面页由 DeckIR.title + subtitle 隐式生成。
SLIDE_CONTENT = "content"
SLIDE_SECTION = "section"
SLIDE_END = "end"
LAYOUT_COVER = "cover"

_TITLE_PLACEHOLDER_TYPES = frozenset(
    {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
        PP_PLACEHOLDER.SUBTITLE,
        PP_PLACEHOLDER.VERTICAL_TITLE,
    }
)
_BODY_PLACEHOLDER_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)

_A_R = qn("a:r")
_A_T = qn("a:t")
_A_RPR = qn("a:rPr")
_A_PPR = qn("a:pPr")
_A_ENDPARARPR = qn("a:endParaRPr")

# OOXML 默认吞掉 <a:t> 首尾空白，需要显式声明保留
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# <a:rPr> / <a:defRPr> 里字体槽的 schema 顺序：latin → ea → cs → sym → …。
# 手工插入 a:ea / a:cs 时必须按这个后继表定位，否则文件在严格校验下不合法。
_SLOT_SUCCESSORS = ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")
_FONT_SLOT_ORDER: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("a:latin", ("a:ea", "a:cs") + _SLOT_SUCCESSORS),
    ("a:ea", ("a:cs",) + _SLOT_SUCCESSORS),
    ("a:cs", _SLOT_SUCCESSORS),
)


# --------------------------------------------------------------------------- #
# DeckIR：模板渲染前的页序与内容
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class Bullet:
    """正文里的一条要点。``level`` 是项目符号层级（0 起）。"""

    text: str
    level: int = 0


@dataclass(frozen=True)
class Slide:
    """DeckIR 的一页：标题 / 要点 / 备注。

    ``kind`` 决定用模板族的哪套版式：``content``（标题 + 正文）、``section``
    （章节页）、``end``（结束页）。封面页不在这里表达，由 ``DeckIR.title`` 与
    ``DeckIR.subtitle`` 隐式生成。
    """

    title: str
    bullets: tuple[Bullet, ...] = ()
    notes: str = ""
    kind: str = SLIDE_CONTENT


@dataclass(frozen=True)
class DeckIR:
    """一份 deck 的页序与内容，即 CJKLayoutPass 的输入（plan §2）。"""

    title: str
    subtitle: str = ""
    slides: tuple[Slide, ...] = ()


@dataclass(frozen=True)
class TemplateSpec:
    """模板族的版式映射：页角色 → 内置母版的 ``slide_layout`` 序号。

    templates/*.json 的“母版 / 占位符”部分加载成这个结构；本层只消费映射结果，
    不解析 JSON。
    """

    name: str
    layouts: Mapping[str, int]

    def layout_index(self, kind: str) -> int:
        try:
            return self.layouts[kind]
        except KeyError:
            raise ValueError(
                f"模板 {self.name!r} 没有 {kind!r} 版式，可选：{sorted(self.layouts)}"
            ) from None


# python-pptx 内置母版：0 封面 / 1 标题和内容 / 2 章节标题 / 5 仅标题。
# 没有模板参数时的默认版式，占位符全部来自真实母版。
DEFAULT_TEMPLATE = TemplateSpec(
    name="default",
    layouts={
        LAYOUT_COVER: 0,
        SLIDE_CONTENT: 1,
        SLIDE_SECTION: 2,
        SLIDE_END: 5,
    },
)


# --------------------------------------------------------------------------- #
# 字体槽与段落属性的读写原语
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class FontSlots:
    """OOXML 的三字体槽：``<a:latin>`` / ``<a:ea>`` / ``<a:cs>``。

    字段为 None 表示该槽不写（保留母版继承）。
    """

    latin: str | None = None
    ea: str | None = None
    cs: str | None = None


@dataclass(frozen=True)
class ParagraphLayout:
    """段落属性快照：``paiban fix`` 的属性 diff 与 ``paiban verify`` 的 lint 都读它。

    ``line_spacing`` 是倍数（``<a:spcPct>``），``line_spacing_pt`` 是固定行距
    （``<a:spcPts>``），两者只会有一个非空。``hanging_punct`` / ``ea_line_break``
    为 None 表示段落上没写这个属性（走母版继承）。
    """

    indent_emu: int | None = None
    line_spacing: float | None = None
    line_spacing_pt: float | None = None
    space_before_pt: float | None = None
    space_after_pt: float | None = None
    alignment: str | None = None
    hanging_punct: bool | None = None
    ea_line_break: bool | None = None


def points_to_emu(points: float) -> int:
    """点 → EMU（DrawingML 的长度单位）。"""
    return int(round(Pt(points).emu))


def read_font_slots(rPr: CT_TextCharacterProperties | None) -> FontSlots | None:
    """读取三字体槽；三个槽都没写时返回 None（字体完全由母版继承）。"""
    if rPr is None:
        return None
    slots = [rPr.find(qn(tag)) for tag, _ in _FONT_SLOT_ORDER]
    if all(slot is None for slot in slots):
        return None
    return FontSlots(*[slot.get("typeface") if slot is not None else None for slot in slots])


def write_font_slots(rPr: CT_TextCharacterProperties, slots: FontSlots) -> None:
    """写三字体槽；值为 None 的槽不动。"""
    for (tag, successors), value in zip(_FONT_SLOT_ORDER, (slots.latin, slots.ea, slots.cs)):
        if value is None:
            continue
        _get_or_add_font_slot(rPr, tag, successors).set("typeface", value)


def _get_or_add_font_slot(rPr: CT_TextCharacterProperties, tag: str, successors: tuple[str, ...]):
    slot = rPr.find(qn(tag))
    if slot is None:
        slot = rPr.makeelement(qn(tag), {})
        rPr.insert_element_before(slot, *successors)
    return slot


def read_paragraph_layout(pPr: CT_TextParagraphProperties | None) -> ParagraphLayout:
    """读取段落属性快照。"""
    if pPr is None:
        return ParagraphLayout()
    indent = pPr.get("indent")
    line_spacing = line_spacing_pt = None
    lnSpc = pPr.lnSpc
    if lnSpc is not None:
        if lnSpc.spcPct is not None:
            line_spacing = lnSpc.spcPct.val
        elif lnSpc.spcPts is not None:
            line_spacing_pt = Emu(lnSpc.spcPts.val).pt
    before = pPr.space_before
    after = pPr.space_after
    return ParagraphLayout(
        indent_emu=int(indent) if indent is not None else None,
        line_spacing=line_spacing,
        line_spacing_pt=line_spacing_pt,
        space_before_pt=Emu(before).pt if before is not None else None,
        space_after_pt=Emu(after).pt if after is not None else None,
        alignment=pPr.get("algn"),
        hanging_punct=_boolean_attribute(pPr, "hangingPunct"),
        ea_line_break=_boolean_attribute(pPr, "eaLnBrk"),
    )


def write_paragraph_layout(
    pPr: CT_TextParagraphProperties,
    *,
    indent_emu: int | None = None,
    line_spacing: float | None = None,
    space_before_pt: float | None = None,
    space_after_pt: float | None = None,
    alignment: str | None = None,
    hanging_punct: bool | None = None,
    ea_line_break: bool | None = None,
) -> None:
    """写段落属性；参数为 None 表示不动该属性。

    ``indent_emu`` 是首行相对段左缩进的偏移（DrawingML 没有“字符”单位，调用方
    按字号换算）；``hanging_punct`` 允许行尾标点悬挂；``ea_line_break`` 打开
    东亚换行与禁则（避头尾）。
    """
    if indent_emu is not None:
        pPr.set("indent", str(int(indent_emu)))
    if line_spacing is not None:
        pPr.line_spacing = float(line_spacing)
    if space_before_pt is not None:
        pPr.space_before = Pt(space_before_pt)
    if space_after_pt is not None:
        pPr.space_after = Pt(space_after_pt)
    if alignment is not None:
        pPr.algn = PP_PARAGRAPH_ALIGNMENT.from_xml(alignment)
    if hanging_punct is not None:
        pPr.set("hangingPunct", "1" if hanging_punct else "0")
    if ea_line_break is not None:
        pPr.set("eaLnBrk", "1" if ea_line_break else "0")


def _boolean_attribute(element, name: str) -> bool | None:
    value = element.get(name)
    if value is None:
        return None
    return value in ("1", "true")


# --------------------------------------------------------------------------- #
# run / 段落 文本原语
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class RunSpan:
    """run 在段落全文里的区间 ``[start, end)``。"""

    run: CT_RegularTextRun
    start: int
    end: int


@dataclass(frozen=True)
class ParagraphText:
    """段落全文 + 每个 run 的区间。

    非 run 的内容子元素（``<a:br>`` / ``<a:fld>``）在全文里记为一个换行符，用作
    相邻性屏障：跨断行的两个标点不算连续标点。
    """

    text: str
    spans: tuple[RunSpan, ...]


def paragraph_runs(paragraph: CT_TextParagraph) -> list[CT_RegularTextRun]:
    """段落里的 ``<a:r>`` 列表（不含 ``<a:br>`` / ``<a:fld>``）。"""
    return list(paragraph.findall(_A_R))


def run_text(run: CT_RegularTextRun) -> str:
    """run 的文本；空 run 返回空串。"""
    text = run.find(_A_T)
    return (text.text or "") if text is not None else ""


def run_properties(run: CT_RegularTextRun) -> CT_TextCharacterProperties:
    """run 的 ``<a:rPr>``，没有就建一个（位置符合 schema 顺序）。"""
    return run.get_or_add_rPr()


def run_font_size_pt(run: CT_RegularTextRun) -> float | None:
    """run 显式字号（pt）；没写返回 None。"""
    rPr = run.find(_A_RPR)
    if rPr is None:
        return None
    return _size_pt(rPr)


def write_run_font_size(rPr: CT_TextCharacterProperties, size_pt: float) -> None:
    rPr.set("sz", str(int(round(size_pt * 100))))


def write_character_spacing(rPr: CT_TextCharacterProperties, points: float) -> None:
    """写 ``<a:rPr spc>``（单位 1/100 pt）；负值即标点挤压。"""
    rPr.set("spc", str(int(round(points * 100))))


def paragraph_font_size_pt(paragraph: CT_TextParagraph) -> float | None:
    """段落字号：取 run 里的最大值，其次段落 ``<a:defRPr>``；都没有返回 None。"""
    sizes = [
        size
        for size in (run_font_size_pt(run) for run in paragraph_runs(paragraph))
        if size is not None
    ]
    if sizes:
        return max(sizes)
    pPr = paragraph.pPr
    if pPr is None or pPr.defRPr is None:
        return None
    return _size_pt(pPr.defRPr)


def paragraph_text_with_runs(paragraph: CT_TextParagraph) -> ParagraphText:
    """段落全文与各 run 的区间；非 run 的内容元素记为换行符屏障。"""
    chunks: list[str] = []
    spans: list[RunSpan] = []
    offset = 0
    for child in paragraph:
        if child.tag in (_A_PPR, _A_ENDPARARPR):
            continue
        if child.tag == _A_R:
            text = run_text(child)
            spans.append(RunSpan(child, offset, offset + len(text)))
            chunks.append(text)
            offset += len(text)
        else:
            # <a:br> / <a:fld> 等：只留下屏障，不参与标点相邻性判断
            chunks.append("\n")
            offset += 1
    return ParagraphText("".join(chunks), tuple(spans))


def split_run(run: CT_RegularTextRun, parts: Sequence[str]) -> list[CT_RegularTextRun]:
    """把 run 按 ``parts`` 拆成多个 run（各自复制原 ``<a:rPr>``）。

    空段会被跳过；``parts`` 拼接后必须与原 run 文本一致，否则抛 ValueError。
    返回的新 run 顺序与 ``parts`` 中非空项一致。
    """
    if "".join(parts) != run_text(run):
        raise ValueError("split_run 的 parts 拼接后必须与原 run 文本一致")
    parent = run.getparent()
    if parent is None:
        raise ValueError("run 不在段落里，无法拆分")
    index = list(parent).index(run)
    new_runs: list[CT_RegularTextRun] = []
    for part in parts:
        if not part:
            continue
        new_run = copy.deepcopy(run)
        set_run_text(new_run, part)
        new_runs.append(new_run)
    parent.remove(run)
    for offset, new_run in enumerate(new_runs):
        parent.insert(index + offset, new_run)
    return new_runs


def set_run_text(run: CT_RegularTextRun, text: str) -> None:
    """写 run 文本；首尾有空白时补 ``xml:space="preserve"``，否则会被读方吞掉。"""
    run.text = text
    if text.strip() != text:
        run.find(_A_T).set(_XML_SPACE, "preserve")


def _size_pt(rPr: CT_TextCharacterProperties) -> float | None:
    size = rPr.get("sz")
    return int(size) / 100 if size is not None else None


# --------------------------------------------------------------------------- #
# 已打开的 .pptx
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class TextBody:
    """一个文本容器：占位符 / 文本框 / 表格单元格 / 备注。"""

    role: str
    label: str
    paragraphs: tuple[CT_TextParagraph, ...]


class Pptx:
    """一个已打开的 .pptx 包，读写 OOXML 的统一入口。"""

    def __init__(self, presentation: Presentation) -> None:
        self._presentation = presentation

    @classmethod
    def open(cls, path: str | Path) -> Pptx:
        return cls(Presentation(str(path)))

    @classmethod
    def create(cls) -> Pptx:
        """空白 deck，用 python-pptx 内置母版（真实占位符，非图片）。"""
        return cls(Presentation())

    @property
    def presentation(self) -> Presentation:
        return self._presentation

    def save(self, path: str | Path) -> None:
        self._presentation.save(str(path))

    def iter_text_bodies(self) -> Iterator[TextBody]:
        """遍历所有文本容器，按 slide 顺序。"""
        for index, slide in enumerate(self._presentation.slides, start=1):
            label = f"slide{index}"
            yield from _iter_shape_bodies(slide.shapes, label)
            if slide.has_notes_slide:
                text_frame = slide.notes_slide.notes_text_frame
                if text_frame is not None:
                    paragraphs = _paragraphs_of(text_frame)
                    if paragraphs:
                        yield TextBody(
                            role=TEXT_ROLE_BODY, label=f"{label}:notes", paragraphs=paragraphs
                        )


def _iter_shape_bodies(shapes, label: str) -> Iterator[TextBody]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _iter_shape_bodies(shape.shapes, label)
        elif shape.has_text_frame:
            paragraphs = _paragraphs_of(shape.text_frame)
            if paragraphs:
                yield TextBody(
                    role=_text_role(shape), label=f"{label}:{shape.name}", paragraphs=paragraphs
                )
        elif shape.has_table:
            for row_index, row in enumerate(shape.table.rows):
                for cell_index, cell in enumerate(row.cells):
                    paragraphs = _paragraphs_of(cell.text_frame)
                    if paragraphs:
                        yield TextBody(
                            role=TEXT_ROLE_BODY,
                            label=f"{label}:{shape.name}[{row_index},{cell_index}]",
                            paragraphs=paragraphs,
                        )


def _paragraphs_of(text_frame) -> tuple[CT_TextParagraph, ...]:
    return tuple(paragraph._p for paragraph in text_frame.paragraphs)


def _text_role(shape) -> str:
    if shape.is_placeholder and shape.placeholder_format.type in _TITLE_PLACEHOLDER_TYPES:
        return TEXT_ROLE_TITLE
    return TEXT_ROLE_BODY


# --------------------------------------------------------------------------- #
# DeckIR → .pptx
# --------------------------------------------------------------------------- #


def emit(deck: DeckIR, template: TemplateSpec = DEFAULT_TEMPLATE) -> Pptx:
    """DeckIR → .pptx：真实母版占位符 + 文本框，无图片兜底。

    封面页来自 ``deck.title`` / ``deck.subtitle``，随后按 ``deck.slides`` 的页序
    依次出页。
    """
    pptx = Pptx.create()
    presentation = pptx.presentation
    _add_slide(
        presentation,
        template.layout_index(LAYOUT_COVER),
        title=deck.title,
        subtitle=deck.subtitle,
    )
    for slide in deck.slides:
        _add_slide(
            presentation,
            template.layout_index(slide.kind),
            title=slide.title,
            bullets=slide.bullets,
            notes=slide.notes,
        )
    return pptx


def _add_slide(
    presentation: Presentation,
    layout_index: int,
    *,
    title: str,
    subtitle: str = "",
    bullets: Sequence[Bullet] = (),
    notes: str = "",
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[layout_index])
    if title:
        title_placeholder = slide.shapes.title
        if title_placeholder is not None:
            title_placeholder.text_frame.text = title
    if subtitle:
        subtitle_placeholder = _find_placeholder(slide, (PP_PLACEHOLDER.SUBTITLE,))
        if subtitle_placeholder is not None:
            subtitle_placeholder.text_frame.text = subtitle
    if bullets:
        body_placeholder = _find_placeholder(slide, _BODY_PLACEHOLDER_TYPES)
        if body_placeholder is not None:
            text_frame = body_placeholder.text_frame
            text_frame.word_wrap = True
            text_frame.text = bullets[0].text
            text_frame.paragraphs[0].level = bullets[0].level
            for bullet in bullets[1:]:
                paragraph = text_frame.add_paragraph()
                paragraph.text = bullet.text
                paragraph.level = bullet.level
    if notes.strip():
        notes_text_frame = slide.notes_slide.notes_text_frame
        if notes_text_frame is not None:
            notes_text_frame.text = notes


def _find_placeholder(slide, types: tuple[PP_PLACEHOLDER, ...]):
    for placeholder in slide.placeholders:
        if placeholder.placeholder_format.type in types:
            return placeholder
    return None
